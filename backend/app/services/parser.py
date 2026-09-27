"""
Extracts plain text from an uploaded resume file (PDF or DOCX).
This text is what gets sent to the AI analysis engine, so the goal
is clean, readable text rather than perfect layout preservation.
"""
import io
from fastapi import HTTPException

import pdfplumber
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read PDF: {exc}")

    text = "\n".join(text_parts).strip()
    if not text:
        raise HTTPException(
            status_code=422,
            detail="No extractable text found in this PDF. It may be a scanned image — try a text-based PDF.",
        )
    return text


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read DOCX: {exc}")

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Include table cell text too (many resumes use tables for layout)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = "\n".join(parts).strip()
    if not text:
        raise HTTPException(status_code=422, detail="No extractable text found in this DOCX file.")
    return text


def extract_text(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Upload a PDF or DOCX.")
